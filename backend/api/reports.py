"""
Génération de rapports Excel — IAMGPS (espace Administrateur).

Un seul classeur .xlsx avec 3 feuilles :
  1. Historique des pointages   (filtrable par module / période)
  2. Statistiques par étudiant  (taux de présence par module)
  3. Alertes de fraude
"""
import datetime
import io

from django.http import HttpResponse
from django.utils import timezone
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import AllowAny

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from .models import Pointage, AlerteFraude, Module, Inscription
from .views import RETARD_TOLERANCE_MIN


ENTETE_FILL = PatternFill(start_color='475569', end_color='475569', fill_type='solid')
ENTETE_FONT = Font(color='FFFFFF', bold=True)


def _ecrire_entetes(feuille, colonnes):
    for i, libelle in enumerate(colonnes, start=1):
        cellule = feuille.cell(row=1, column=i, value=libelle)
        cellule.fill = ENTETE_FILL
        cellule.font = ENTETE_FONT
        cellule.alignment = Alignment(horizontal='center')
    feuille.freeze_panes = 'A2'


def _ajuster_largeur_colonnes(feuille):
    for colonne in feuille.columns:
        longueur_max = 0
        lettre = get_column_letter(colonne[0].column)
        for cellule in colonne:
            valeur = str(cellule.value) if cellule.value is not None else ''
            longueur_max = max(longueur_max, len(valeur))
        feuille.column_dimensions[lettre].width = min(longueur_max + 3, 45)


@api_view(['GET'])
@permission_classes([AllowAny])
def export_rapport_excel(request):
    """
    GET /api/rapports/export/

    Paramètres optionnels (query string) :
      - module      : code du module (ex: SYC-BI) — sinon tous les modules
      - date_debut  : YYYY-MM-DD — sinon pas de borne basse
      - date_fin    : YYYY-MM-DD — sinon pas de borne haute

    Retourne un fichier .xlsx en téléchargement direct.
    """
    module_code = request.GET.get('module')
    date_debut = request.GET.get('date_debut')
    date_fin = request.GET.get('date_fin')

    pointages = Pointage.objects.select_related(
        'utilisateur', 'session', 'session__module', 'session__zone'
    ).order_by('session__date', 'timestamp')

    if module_code:
        pointages = pointages.filter(session__module__code=module_code)
    if date_debut:
        pointages = pointages.filter(session__date__gte=date_debut)
    if date_fin:
        pointages = pointages.filter(session__date__lte=date_fin)

    wb = Workbook()

    # ── Feuille 1 : Historique des pointages ──────────────────────
    ws1 = wb.active
    ws1.title = 'Historique pointages'
    _ecrire_entetes(ws1, [
        'Date', 'Heure', 'Module', 'Étudiant', 'Matricule', 'Statut',
        'Distance zone (m)', 'Précision GPS (m)', 'Latitude', 'Longitude',
    ])

    ligne = 2
    for p in pointages:
        heure_locale = timezone.localtime(p.timestamp)
        ws1.cell(row=ligne, column=1, value=p.session.date.strftime('%d/%m/%Y'))
        ws1.cell(row=ligne, column=2, value=heure_locale.strftime('%H:%M:%S'))
        ws1.cell(row=ligne, column=3, value=p.session.module.code)
        ws1.cell(row=ligne, column=4, value=f"{p.utilisateur.prenom} {p.utilisateur.nom}")
        ws1.cell(row=ligne, column=5, value=p.utilisateur.matricule)
        ws1.cell(row=ligne, column=6, value=p.get_statut_display())
        ws1.cell(row=ligne, column=7, value=round(p.distance_zone_m, 1) if p.distance_zone_m is not None else '')
        ws1.cell(row=ligne, column=8, value=round(p.precision_gps_m, 1) if p.precision_gps_m is not None else '')
        ws1.cell(row=ligne, column=9, value=float(p.latitude))
        ws1.cell(row=ligne, column=10, value=float(p.longitude))
        ligne += 1
    _ajuster_largeur_colonnes(ws1)

    # ── Feuille 2 : Statistiques par étudiant ─────────────────────
    ws2 = wb.create_sheet('Statistiques étudiants')
    _ecrire_entetes(ws2, [
        'Matricule', 'Étudiant', 'Module', 'Séances terminées',
        'Présences', 'Retards', 'Absences', 'Taux (%)',
    ])

    modules_qs = Module.objects.all()
    if module_code:
        modules_qs = modules_qs.filter(code=module_code)

    ligne = 2
    for module in modules_qs:
        sessions_terminees = module.sessions.filter(statut='TERMINE')
        if date_debut:
            sessions_terminees = sessions_terminees.filter(date__gte=date_debut)
        if date_fin:
            sessions_terminees = sessions_terminees.filter(date__lte=date_fin)
        nb_sessions = sessions_terminees.count()
        if nb_sessions == 0:
            continue

        inscriptions = Inscription.objects.filter(module=module, is_active=True).select_related('etudiant')
        for insc in inscriptions:
            etu = insc.etudiant
            pointages_etu = Pointage.objects.filter(
                utilisateur=etu, session__in=sessions_terminees, statut='VALIDE'
            ).select_related('session')

            nb_retard, nb_present = 0, 0
            for p in pointages_etu:
                limite = (
                    datetime.datetime.combine(p.session.date, p.session.heure_debut)
                    + datetime.timedelta(minutes=RETARD_TOLERANCE_MIN)
                ).time()
                if timezone.localtime(p.timestamp).time() > limite:
                    nb_retard += 1
                else:
                    nb_present += 1
            nb_absent = max(0, nb_sessions - nb_present - nb_retard)
            taux = round((nb_present + nb_retard) / nb_sessions * 100, 1) if nb_sessions else 0

            ws2.cell(row=ligne, column=1, value=etu.matricule)
            ws2.cell(row=ligne, column=2, value=f"{etu.prenom} {etu.nom}")
            ws2.cell(row=ligne, column=3, value=module.code)
            ws2.cell(row=ligne, column=4, value=nb_sessions)
            ws2.cell(row=ligne, column=5, value=nb_present)
            ws2.cell(row=ligne, column=6, value=nb_retard)
            ws2.cell(row=ligne, column=7, value=nb_absent)
            ws2.cell(row=ligne, column=8, value=taux)
            ligne += 1
    _ajuster_largeur_colonnes(ws2)

    # ── Feuille 3 : Alertes de fraude ──────────────────────────────
    ws3 = wb.create_sheet('Alertes de fraude')
    _ecrire_entetes(ws3, [
        'Date', 'Étudiant', 'Matricule', 'Module', 'Type de fraude',
        'Score IA (%)', 'Niveau', 'Statut', 'Traité par', 'Commentaire',
    ])

    alertes = AlerteFraude.objects.select_related(
        'pointage__utilisateur', 'pointage__session__module', 'traite_par'
    ).order_by('-timestamp')
    if module_code:
        alertes = alertes.filter(pointage__session__module__code=module_code)
    if date_debut:
        alertes = alertes.filter(pointage__session__date__gte=date_debut)
    if date_fin:
        alertes = alertes.filter(pointage__session__date__lte=date_fin)

    ligne = 2
    for a in alertes:
        etu = a.pointage.utilisateur
        ws3.cell(row=ligne, column=1, value=timezone.localtime(a.timestamp).strftime('%d/%m/%Y %H:%M'))
        ws3.cell(row=ligne, column=2, value=f"{etu.prenom} {etu.nom}")
        ws3.cell(row=ligne, column=3, value=etu.matricule)
        ws3.cell(row=ligne, column=4, value=a.pointage.session.module.code)
        ws3.cell(row=ligne, column=5, value=a.get_type_fraude_display())
        ws3.cell(row=ligne, column=6, value=round(a.score_ia, 1) if a.score_ia is not None else '')
        ws3.cell(row=ligne, column=7, value=a.get_niveau_severite_display() if a.niveau_severite else '')
        ws3.cell(row=ligne, column=8, value=a.get_statut_display())
        ws3.cell(row=ligne, column=9, value=f"{a.traite_par.prenom} {a.traite_par.nom}" if a.traite_par else '')
        ws3.cell(row=ligne, column=10, value=a.commentaire or '')
        ligne += 1
    _ajuster_largeur_colonnes(ws3)

    # ── Réponse HTTP : fichier à télécharger ───────────────────────
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nom_fichier = f"rapport_iamgps_{timezone.localdate().strftime('%Y%m%d')}.xlsx"
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{nom_fichier}"'
    return response